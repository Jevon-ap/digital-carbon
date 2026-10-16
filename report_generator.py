"""
Modul Generator Laporan Emisi GRK (PDF, Excel, CSV)
Kalkulator Hijau V.2 (2026)
"""

import io
from datetime import datetime
from typing import Dict, Any, Optional
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas


class NumberedCanvas(canvas.Canvas):
    """Canvas kustom untuk menambahkan nomor halaman otomatis 'Halaman X dari Y'."""
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#555555"))
        footer_text = f"Kalkulator Hijau V.2 (2026) | Standar GHG Protocol & PCAF | Halaman {self._pageNumber} dari {page_count}"
        self.drawRightString(A4[0] - 40, 25, footer_text)
        self.drawString(40, 25, f"Dicetak: {datetime.now().strftime('%d %B %Y %H:%M WIB')}")
        self.setStrokeColor(colors.HexColor("#D0D5DD"))
        self.setLineWidth(0.5)
        self.line(40, 36, A4[0] - 40, 36)
        self.restoreState()


def generate_excel_report(
    summary: Any,
    df_s1: pd.DataFrame,
    df_s2: pd.DataFrame,
    df_s3: pd.DataFrame,
    df_s4: pd.DataFrame,
    company_info: Optional[Dict[str, str]] = None
) -> bytes:
    """Membuat laporan emisi multi-sheet Excel yang diformat profesional."""
    wb = Workbook()
    info = company_info or {
        "entitas": "Entitas Pelapor / Lembaga Keuangan",
        "tahun": "2026",
        "penilai": "Tim Keberlanjutan & ESG"
    }

    # Style Palet Hijau Profesional
    header_fill = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
    sub_fill = PatternFill(start_color="2C6B56", end_color="2C6B56", fill_type="solid")
    kpi_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    normal_font = Font(name="Segoe UI", size=10)
    title_font = Font(name="Segoe UI", size=14, bold=True, color="1B4D3E")
    
    thin_border = Border(
        left=Side(style='thin', color='D0D5DD'),
        right=Side(style='thin', color='D0D5DD'),
        top=Side(style='thin', color='D0D5DD'),
        bottom=Side(style='thin', color='D0D5DD')
    )

    # 1. Sheet Ringkasan Eksekutif
    ws_sum = wb.active
    ws_sum.title = "Ringkasan Eksekutif"
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum["A1"] = "LAPORAN INVENTARISASI EMISI GAS RUMAH KACA (GRK)"
    ws_sum["A1"].font = title_font
    ws_sum["A2"] = f"Buku Panduan Kalkulator Hijau V.2 (2026) | Entitas: {info.get('entitas', '')} | Tahun: {info.get('tahun', '2026')}"
    ws_sum["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="555555")

    ws_sum.append([])
    ws_sum.append(["METRIK EMISI", "NILAI (tCO2e)", "KETERANGAN & STATUS"])
    for col in range(1, 4):
        c = ws_sum.cell(row=4, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    summary_rows = [
        ("Scope 1 (Emisi Langsung - BBM & Kendaraan)", summary.scope1_tco2e, f"{summary.count_scope1} aktivitas"),
        ("Scope 2 (Emisi Listrik Jaringan / PLN)", summary.scope2_tco2e, f"{summary.count_scope2} fasilitas"),
        ("Scope 3 (Financed Emissions - PCAF)", summary.scope3_tco2e, f"{summary.count_scope3} fasilitas kredit/debitur"),
        ("TOTAL GROSS EMISSIONS (Scope 1 + 2 + 3)", summary.gross_tco2e, "Total Emisi Kotor"),
        ("Scope 4 (Avoided Emissions & Offsetting)", summary.scope4_tco2e, f"{summary.count_scope4} inisiatif hijau / SPE-GRK"),
        ("TOTAL NET EMISSIONS (Gross - Scope 4)", summary.net_tco2e, "Total Emisi Bersih Terkoreksi"),
        ("Persentase Penurunan Emisi (Offsets %)", f"{summary.offset_percentage:.2f}%", "Rasio Penurunan terhadap Emisi Kotor")
    ]

    for idx, (m, v, k) in enumerate(summary_rows, start=5):
        ws_sum.cell(row=idx, column=1, value=m).font = bold_font if "TOTAL" in m else normal_font
        ws_sum.cell(row=idx, column=2, value=v).font = bold_font if "TOTAL" in m else normal_font
        ws_sum.cell(row=idx, column=3, value=k).font = normal_font
        if "TOTAL" in m:
            for c in range(1, 4):
                ws_sum.cell(row=idx, column=c).fill = kpi_fill
        for c in range(1, 4):
            ws_sum.cell(row=idx, column=c).border = thin_border

    # Fungsi bantu untuk menulis DataFrame ke sheet
    def write_df_to_sheet(sheet, title, df, scope_tag):
        sheet.views.sheetView[0].showGridLines = True
        sheet["A1"] = f"{scope_tag} - {title}"
        sheet["A1"].font = title_font
        sheet.append([])
        if df.empty:
            sheet.append(["Tidak ada data tercatat."])
            return
        
        headers = list(df.columns)
        sheet.append([h.replace("_", " ").title() for h in headers])
        header_row = 3
        for col_idx, _ in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=col_idx)
            cell.fill = sub_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for _, row in df.iterrows():
            sheet.append(list(row.values))
            
        for r in range(header_row + 1, sheet.max_row + 1):
            for c in range(1, len(headers) + 1):
                cell = sheet.cell(row=r, column=c)
                cell.font = normal_font
                cell.border = thin_border

    # 2. Scope 1 Sheet
    ws_s1 = wb.create_sheet(title="Scope 1")
    write_df_to_sheet(ws_s1, "Emisi Langsung BBM & Kendaraan", df_s1, "SCOPE 1")

    # 3. Scope 2 Sheet
    ws_s2 = wb.create_sheet(title="Scope 2")
    write_df_to_sheet(ws_s2, "Emisi Tidak Langsung Pembelian Listrik", df_s2, "SCOPE 2")

    # 4. Scope 3 Sheet
    ws_s3 = wb.create_sheet(title="Scope 3")
    write_df_to_sheet(ws_s3, "Financed Emissions Sektor Keuangan (PCAF)", df_s3, "SCOPE 3")

    # 5. Scope 4 Sheet
    ws_s4 = wb.create_sheet(title="Scope 4")
    write_df_to_sheet(ws_s4, "Avoided Emissions & Karbon Offsetting", df_s4, "SCOPE 4")

    # Atur lebar kolom otomatis untuk seluruh sheet
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_report(
    summary: Any,
    df_s1: pd.DataFrame,
    df_s2: pd.DataFrame,
    df_s3: pd.DataFrame,
    df_s4: pd.DataFrame,
    company_info: Optional[Dict[str, str]] = None
) -> bytes:
    """Membuat laporan emisi GRK profesional format PDF menggunakan ReportLab."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()
    
    # Kustomisasi styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1B4D3E")
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#4A5568")
    )
    h2_style = ParagraphStyle(
        "SectionH2",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1B4D3E"),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        "BodySmall",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#2D3748")
    )
    table_cell = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1A202C")
    )
    table_cell_bold = ParagraphStyle(
        "TableCellBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1A202C")
    )
    table_header = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white
    )

    info = company_info or {
        "entitas": "Entitas Pelapor",
        "tahun": "2026",
        "penilai": "Tim Keberlanjutan & ESG"
    }

    elements = []

    # Header / Banner
    elements.append(Paragraph("LAPORAN INVENTARISASI EMISI GAS RUMAH KACA (GRK)", title_style))
    elements.append(Spacer(1, 2))
    elements.append(Paragraph(
        f"<b>Standar Perhitungan:</b> Buku Panduan Kalkulator Hijau V.2 (2026) | GHG Protocol | PCAF Standard<br/>"
        f"<b>Entitas:</b> {info.get('entitas', '-')} | <b>Periode:</b> Tahun {info.get('tahun', '2026')} | <b>Penilai:</b> {info.get('penilai', '-')}",
        subtitle_style
    ))
    elements.append(Spacer(1, 6))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1B4D3E"), spaceAfter=10))

    # Ringkasan Eksekutif - KPI Cards Table
    elements.append(Paragraph("1. RINGKASAN EKSEKUTIF INVENTARISASI EMISI", h2_style))
    
    kpi_data = [
        [
            Paragraph("<b>TOTAL EMISI KOTOR (GROSS)</b><br/><font size=14 color='#C53030'><b>" f"{summary.gross_tco2e:,.2f}" "</b></font> tCO2e", body_style),
            Paragraph("<b>PENGURANGAN & OFFSET (SCOPE 4)</b><br/><font size=14 color='#2B6CB0'><b>" f"{summary.scope4_tco2e:,.2f}" "</b></font> tCO2e", body_style),
            Paragraph("<b>TOTAL EMISI BERSIH (NET)</b><br/><font size=14 color='#276749'><b>" f"{summary.net_tco2e:,.2f}" "</b></font> tCO2e", body_style),
            Paragraph("<b>OFFSET RATIO</b><br/><font size=14 color='#234E52'><b>" f"{summary.offset_percentage:.1f}%" "</b></font>", body_style)
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[130, 130, 130, 133])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F7FAFC")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#CBD5E0")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 10))

    # Tabel Rincian per Scope
    elements.append(Paragraph("2. KOMPOSISI EMISI PER LINGKUP (SCOPE)", h2_style))
    summary_table_data = [
        [
            Paragraph("Lingkup Emisi (Scope)", table_header),
            Paragraph("Jumlah Aktivitas", table_header),
            Paragraph("Total Emisi (tCO2e)", table_header),
            Paragraph("Porsi Gross (%)", table_header),
            Paragraph("Keterangan Metodologi", table_header)
        ],
        [
            Paragraph("<b>Scope 1</b> (Emisi Langsung)", table_cell_bold),
            Paragraph(str(summary.count_scope1), table_cell),
            Paragraph(f"{summary.scope1_tco2e:,.4f}", table_cell),
            Paragraph(f"{(summary.scope1_tco2e/summary.gross_tco2e*100 if summary.gross_tco2e>0 else 0):.1f}%", table_cell),
            Paragraph("Pembakaran BBM stasioner & jarak tempuh kendaraan operasional", table_cell)
        ],
        [
            Paragraph("<b>Scope 2</b> (Emisi Listrik Jaringan)", table_cell_bold),
            Paragraph(str(summary.count_scope2), table_cell),
            Paragraph(f"{summary.scope2_tco2e:,.4f}", table_cell),
            Paragraph(f"{(summary.scope2_tco2e/summary.gross_tco2e*100 if summary.gross_tco2e>0 else 0):.1f}%", table_cell),
            Paragraph("Konsumsi listrik kWh x Faktor Emisi Grid Regional/Nasional", table_cell)
        ],
        [
            Paragraph("<b>Scope 3</b> (Financed Emissions)", table_cell_bold),
            Paragraph(str(summary.count_scope3), table_cell),
            Paragraph(f"{summary.scope3_tco2e:,.4f}", table_cell),
            Paragraph(f"{(summary.scope3_tco2e/summary.gross_tco2e*100 if summary.gross_tco2e>0 else 0):.1f}%", table_cell),
            Paragraph("PCAF standard (Reported Debitur & Proksi Sektoral)", table_cell)
        ],
        [
            Paragraph("<b>TOTAL EMISI KOTOR (GROSS)</b>", table_cell_bold),
            Paragraph(str(summary.count_scope1 + summary.count_scope2 + summary.count_scope3), table_cell_bold),
            Paragraph(f"<b>{summary.gross_tco2e:,.4f}</b>", table_cell_bold),
            Paragraph("<b>100.0%</b>", table_cell_bold),
            Paragraph("<b>Total Beban Emisi Sebelum Penurunan</b>", table_cell_bold)
        ],
        [
            Paragraph("<b>Scope 4</b> (Avoided & Offsets)", table_cell_bold),
            Paragraph(str(summary.count_scope4), table_cell),
            Paragraph(f"-{summary.scope4_tco2e:,.4f}", table_cell),
            Paragraph(f"-{summary.offset_percentage:.1f}%", table_cell),
            Paragraph("Penghindaran EBT/EV & Kredit Karbon SPE-GRK", table_cell)
        ],
        [
            Paragraph("<b>TOTAL EMISI BERSIH (NET)</b>", table_cell_bold),
            Paragraph("-", table_cell_bold),
            Paragraph(f"<b>{summary.net_tco2e:,.4f}</b>", table_cell_bold),
            Paragraph("-", table_cell_bold),
            Paragraph("<b>Gross dikurangi Pengurangan Scope 4 Terverifikasi</b>", table_cell_bold)
        ]
    ]

    t_summary = Table(summary_table_data, colWidths=[140, 75, 95, 75, 138])
    t_summary.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1B4D3E")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E0")),
        ('BACKGROUND', (0,4), (-1,4), colors.HexColor("#EDF2F7")),
        ('BACKGROUND', (0,6), (-1,6), colors.HexColor("#E6FFFA")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(t_summary)
    elements.append(Spacer(1, 14))

    # Fungsi bantu untuk membuat tabel preview rincian
    def create_table_preview(df: pd.DataFrame, title: str, cols: list, col_widths: list, value_col: str):
        if df.empty:
            return
        elements.append(Paragraph(title, h2_style))
        header_row = [Paragraph(c, table_header) for c in cols]
        data_rows = [header_row]
        for _, r in df.head(10).iterrows():
            row = []
            for c in cols:
                val = r.get(c.lower().replace(" ", "_"), "-")
                if isinstance(val, (int, float)):
                    val_str = f"{val:,.2f}"
                else:
                    val_str = str(val)[:28]
                row.append(Paragraph(val_str, table_cell))
            data_rows.append(row)
        
        t = Table(data_rows, colWidths=col_widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#2C6B56")),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ]))
        elements.append(t)
        if len(df) > 10:
            elements.append(Paragraph(f"<i>*Menampilkan 10 dari {len(df)} baris. Unduh berkas Excel untuk rincian lengkap.</i>", subtitle_style))
        elements.append(Spacer(1, 8))

    # Preview Rincian Scope 1-4
    if not df_s1.empty:
        cols_s1 = ["Deskripsi", "Jenis BBM", "Sub Kategori", "Emisi Tco2E"]
        create_table_preview(df_s1, "3. RINCIAN DATA SCOPE 1 (EMISI LANGSUNG)", cols_s1, [180, 140, 110, 93], "emisi_tco2e")

    if not df_s2.empty:
        cols_s2 = ["Deskripsi", "Wilayah Grid", "Konsumsi Kwh", "Emisi Tco2E"]
        create_table_preview(df_s2, "4. RINCIAN DATA SCOPE 2 (EMISI LISTRIK)", cols_s2, [180, 150, 100, 93], "emisi_tco2e")

    if not df_s3.empty:
        cols_s3 = ["Nama Debitur", "Metode", "Sektor", "Emisi Tco2E"]
        create_table_preview(df_s3, "5. RINCIAN DATA SCOPE 3 (FINANCED EMISSIONS)", cols_s3, [160, 110, 160, 93], "emisi_tco2e")

    if not df_s4.empty:
        cols_s4 = ["Deskripsi", "Kategori", "Avoided Tco2E"]
        create_table_preview(df_s4, "6. RINCIAN DATA SCOPE 4 (AVOIDED & OFFSETS)", cols_s4, [240, 160, 123], "avoided_tco2e")

    # Catatan Metodologi & Legalitas
    elements.append(KeepTogether([
        Spacer(1, 8),
        HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#CBD5E0"), spaceAfter=6),
        Paragraph("<b>Keterangan Metodologi & Kepatuhan:</b>", h2_style),
        Paragraph(
            "1. Perhitungan Scope 1 dan Scope 2 mengacu pada <i>Pedoman Teknis Perhitungan Emisi GRK Sektor Energi dan Transportasi KLHK / ESDM</i> serta <i>GHG Protocol Corporate Standard</i>.<br/>"
            "2. Perhitungan Scope 3 (Financed Emissions) mengacu pada standar global <i>PCAF (Partnership for Carbon Accounting Financials)</i> dan <i>Taksonomi Hijau Indonesia / Taksonomi Keuangan Berkelanjutan OJK</i>.<br/>"
            "3. Scope 4 dicatat secara terpisah sebagai penurunan emisi (Avoided Emissions & Offsetting) dan tidak mengurangi emisi kotor secara akuntansi langsung, namun dilaporkan sebagai Net Offsetting terverifikasi.",
            body_style
        ),
        Spacer(1, 16),
        Table([
            [
                Paragraph("<b>Disiapkan Oleh:</b><br/><br/><br/>______________________<br/>Analis Keberlanjutan & ESG", body_style),
                Paragraph("<b>Ditinjau & Disetujui Oleh:</b><br/><br/><br/>______________________<br/>Kepala Divisi Risiko & Keberlanjutan", body_style)
            ]
        ], colWidths=[260, 263])
    ]))

    doc.build(elements, canvasmaker=NumberedCanvas)
    return buf.getvalue()


def generate_consolidated_csv(
    summary: Any,
    df_s1: pd.DataFrame,
    df_s2: pd.DataFrame,
    df_s3: pd.DataFrame,
    df_s4: pd.DataFrame
) -> str:
    """Menggabungkan seluruh tabel ke dalam teks format CSV terpadu."""
    lines = []
    lines.append("# ====================================================================")
    lines.append("# LAPORAN INVENTARISASI EMISI GRK - KALKULATOR HIJAU V.2 (2026)")
    lines.append(f"# Dibuat pada: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append("# ====================================================================")
    lines.append("")
    lines.append("# --- RINGKASAN EKSEKUTIF ---")
    lines.append("Lingkup,Emisi_tCO2e,Keterangan")
    lines.append(f"Scope 1,{summary.scope1_tco2e},Emisi Langsung BBM & Kendaraan")
    lines.append(f"Scope 2,{summary.scope2_tco2e},Emisi Pembelian Listrik Jaringan")
    lines.append(f"Scope 3,{summary.scope3_tco2e},Financed Emissions (PCAF)")
    lines.append(f"GROSS EMISSIONS,{summary.gross_tco2e},Total Scope 1+2+3")
    lines.append(f"Scope 4 (Offsets/Avoided),{summary.scope4_tco2e},Pengurangan & Kredit Karbon")
    lines.append(f"NET EMISSIONS,{summary.net_tco2e},Gross - Scope 4")
    lines.append(f"Offset Percentage,{summary.offset_percentage}%,Rasio Penurunan")
    lines.append("")

    for title, df in [
        ("SCOPE 1 - EMISI LANGSUNG", df_s1),
        ("SCOPE 2 - EMISI LISTRIK", df_s2),
        ("SCOPE 3 - FINANCED EMISSIONS", df_s3),
        ("SCOPE 4 - AVOIDED EMISSIONS & OFFSETS", df_s4)
    ]:
        lines.append(f"# --- {title} ---")
        if not df.empty:
            csv_str = df.to_csv(index=False).strip()
            lines.append(csv_str)
        else:
            lines.append("Status: Tidak ada data")
        lines.append("")

    return "\n".join(lines)
